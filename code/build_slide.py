import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
def text(x, y, w, h, t, size, bold=False, color=(0x20, 0x20, 0x20)):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor(*color); return tf
text(0.5, 0.3, 12.3, 0.9, "Aiki-Sol external check: AUROC 0.858 (0.823\u20130.891) on 500 test-split proteins; simple baselines 0.608 and 0.506", 24, True)
def pic(path, x, y, w):
    im = Image.open(path); h = w * im.height / im.width; s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h)); return h
h1 = pic(os.path.join(ROOT, "figures", "fig1_roc_pr.png"), 0.4, 1.55, 7.9)
h2 = pic(os.path.join(ROOT, "figures", "fig2_calibration.png"), 8.7, 1.45, 4.2)
y = 1.55 + max(h1, h2) + 0.2
text(0.5, y, 12.3, 0.5, "Left: ROC and precision\u2013recall, shaded bands = 95% bootstrap intervals. Right: reliability of the marginal output (Brier 0.156).", 14)
text(0.5, y + 0.55, 12.3, 1.0, "Caveat: 64.8% of the 500 test sequences exactly match Aiki-Sol\u2019s published training list, so 0.858 is an upper bound; AUROC is 0.741 (0.668\u20130.807) on the 176 sequences that do not match. One benchmark, N = 500.", 14, True, (0x8a, 0x2b, 0x00))
text(0.5, 7.08, 12.3, 0.3, "Aaron Ezekiel \u00b7 UC Berkeley \u00b7 independent evaluation, October 2026", 11, False, (0x55, 0x55, 0x55))
out = os.path.join(ROOT, "slides"); os.makedirs(out, exist_ok=True); prs.save(os.path.join(out, "aikisol_eval_figures_slide.pptx")); print("saved", h1, h2, y)
